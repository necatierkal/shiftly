import io
from datetime import datetime, date
from typing import Dict, Any, List
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models.models import DutyPlan, DutyPlanShift, Team, DutyPeriod, User


class ExportService:
    TURKISH_DAYS = {
        0: "Pazartesi",
        1: "Salı",
        2: "Çarşamba",
        3: "Perşembe",
        4: "Cuma",
        5: "Cumartesi",
        6: "Pazar",
    }

    @staticmethod
    def generate_excel(plan: DutyPlan, db: Session) -> io.BytesIO:
        """Generates a professional multi-sheet Excel (.xlsx) file for a duty plan."""
        team = plan.team
        period = plan.period
        shifts: List[DutyPlanShift] = (
            db.query(DutyPlanShift)
            .filter(DutyPlanShift.plan_id == plan.id)
            .order_by(DutyPlanShift.date.asc(), DutyPlanShift.start_datetime.asc())
            .all()
        )

        wb = openpyxl.Workbook()

        # Styles
        primary_color = "3157D5"
        light_blue = "E8EDFB"
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="111827")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)
        warning_font = Font(name="Calibri", size=11, bold=True, color="DC2626")

        header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
        zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        warning_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB"),
        )

        # -------------------------------------------------------------
        # SHEET 1: Nöbet Çizelgesi
        # -------------------------------------------------------------
        ws_schedule = wb.active
        ws_schedule.title = "Nöbet Çizelgesi"

        # Title Block
        ws_schedule["A1"] = f"{team.name.upper()} — RESMİ NÖBET ÇİZELGESİ"
        ws_schedule["A1"].font = title_font
        ws_schedule["A2"] = (
            f"Dönem: {period.name} | Sürüm: v{plan.version} | "
            f"Durum: {plan.status} | Yayımlanma: {plan.published_at.strftime('%d.%m.%Y %H:%M') if plan.published_at else 'Taslak'}"
        )
        ws_schedule["A2"].font = subtitle_font

        # Headers
        headers = [
            "Tarih",
            "Gün",
            "Vardiya / Nöbet Türü",
            "Başlangıç",
            "Bitiş",
            "Süre (Saat)",
            "Görevli Personel",
            "Konum",
            "Açıklama / Durum",
        ]

        row_idx = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws_schedule.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
        ws_schedule.row_dimensions[row_idx].height = 28

        # Data Rows
        person_stats: Dict[int, Dict[str, Any]] = {}

        for s in shifts:
            row_idx += 1
            day_name = ExportService.TURKISH_DAYS.get(s.date.weekday(), "")
            duration = (s.end_datetime - s.start_datetime).total_seconds() / 3600.0

            staff_name = s.assigned_user.fullName if s.assigned_user else "ATANMADI"
            status_note = s.unassigned_reason if s.is_unassigned else (s.violation_warning or s.note or "")

            row_data = [
                s.date.strftime("%d.%m.%Y"),
                day_name,
                s.shift_type.name if s.shift_type else "Vardiya",
                s.start_datetime.strftime("%H:%M"),
                s.end_datetime.strftime("%H:%M"),
                duration,
                staff_name,
                s.location or "-",
                status_note,
            ]

            # Aggregate stats for Sheet 2
            if s.assigned_user:
                uid = s.assigned_user.id
                if uid not in person_stats:
                    person_stats[uid] = {
                        "name": s.assigned_user.fullName,
                        "job_title": s.assigned_user.job_title or "-",
                        "total_shifts": 0,
                        "total_hours": 0.0,
                        "night_hours": 0.0,
                        "weekend_shifts": 0,
                    }
                person_stats[uid]["total_shifts"] += 1
                person_stats[uid]["total_hours"] += duration
                if s.shift_type and s.shift_type.is_night:
                    person_stats[uid]["night_hours"] += duration
                if s.date.weekday() in (5, 6):
                    person_stats[uid]["weekend_shifts"] += 1

            for col_idx, val in enumerate(row_data, 1):
                cell = ws_schedule.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.font = warning_font if s.is_unassigned and col_idx == 7 else regular_font

                if row_idx % 2 == 0:
                    cell.fill = zebra_fill
                if s.is_unassigned:
                    cell.fill = warning_fill

                if col_idx in [1, 2, 4, 5, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_schedule.row_dimensions[row_idx].height = 22

        # Auto-fit column widths
        for col in ws_schedule.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws_schedule.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # -------------------------------------------------------------
        # SHEET 2: Personel Nöbet İcmali (Summary)
        # -------------------------------------------------------------
        ws_summary = wb.create_sheet(title="Personel İcmali")
        ws_summary["A1"] = f"{team.name.upper()} — PERSONEL NÖBET DAĞILIM İCMALİ"
        ws_summary["A1"].font = title_font
        ws_summary["A2"] = f"Dönem: {period.name} | Toplam Nöbet Sayısı: {len(shifts)}"
        ws_summary["A2"].font = subtitle_font

        sum_headers = [
            "Personel Adı Soyadı",
            "Görevi / Unvanı",
            "Toplam Nöbet",
            "Toplam Saat",
            "Gece Nöbeti (Saat)",
            "Hafta Sonu Nöbeti",
        ]

        row_idx = 4
        for col_idx, h in enumerate(sum_headers, 1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws_summary.row_dimensions[row_idx].height = 28

        for uid, p in person_stats.items():
            row_idx += 1
            vals = [
                p["name"],
                p["job_title"],
                p["total_shifts"],
                round(p["total_hours"], 1),
                round(p["night_hours"], 1),
                p["weekend_shifts"],
            ]
            for col_idx, v in enumerate(vals, 1):
                cell = ws_summary.cell(row=row_idx, column=col_idx, value=v)
                cell.border = thin_border
                cell.font = regular_font
                if col_idx in [3, 4, 5, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_summary.row_dimensions[row_idx].height = 22

        # Auto-fit Sheet 2
        for col in ws_summary.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_pdf(plan: DutyPlan, db: Session) -> io.BytesIO:
        """Generates a printable, styled PDF document for a duty plan."""
        team = plan.team
        period = plan.period
        shifts: List[DutyPlanShift] = (
            db.query(DutyPlanShift)
            .filter(DutyPlanShift.plan_id == plan.id)
            .order_by(DutyPlanShift.date.asc(), DutyPlanShift.start_datetime.asc())
            .all()
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(A4),
            rightMargin=25,
            leftMargin=25,
            topMargin=25,
            bottomMargin=25,
        )

        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E3A8A"),
            alignment=1,  # Center
            spaceAfter=4,
        )

        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#4B5563"),
            alignment=1,
            spaceAfter=14,
        )

        cell_style = ParagraphStyle(
            "CellText",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#111827"),
        )

        cell_header_style = ParagraphStyle(
            "CellHeader",
            parent=styles["Normal"],
            fontSize=9,
            leading=11,
            textColor=colors.white,
            fontName="Helvetica-Bold",
            alignment=1,
        )

        pub_date_str = plan.published_at.strftime("%d.%m.%Y %H:%M") if plan.published_at else "Taslak Plan"
        elements.append(Paragraph(f"<b>{team.name.upper()} — RESMİ NÖBET ÇİZELGESİ</b>", title_style))
        elements.append(
            Paragraph(
                f"<b>Dönem:</b> {period.name} ({period.start_date.strftime('%d.%m.%Y')} - {period.end_date.strftime('%d.%m.%Y')}) &nbsp;|&nbsp; "
                f"<b>Plan Sürümü:</b> v{plan.version} &nbsp;|&nbsp; "
                f"<b>Durum:</b> {plan.status} &nbsp;|&nbsp; "
                f"<b>Yayımlanma:</b> {pub_date_str}",
                subtitle_style,
            )
        )

        # Table data
        table_data = [
            [
                Paragraph("<b>Tarih</b>", cell_header_style),
                Paragraph("<b>Gün</b>", cell_header_style),
                Paragraph("<b>Nöbet Türü</b>", cell_header_style),
                Paragraph("<b>Saat Aralığı</b>", cell_header_style),
                Paragraph("<b>Süre</b>", cell_header_style),
                Paragraph("<b>Görevli Personel</b>", cell_header_style),
                Paragraph("<b>Konum</b>", cell_header_style),
                Paragraph("<b>Not / Durum</b>", cell_header_style),
            ]
        ]

        for s in shifts:
            day_name = ExportService.TURKISH_DAYS.get(s.date.weekday(), "")
            duration = (s.end_datetime - s.start_datetime).total_seconds() / 3600.0
            time_str = f"{s.start_datetime.strftime('%H:%M')} - {s.end_datetime.strftime('%H:%M')}"
            staff_name = s.assigned_user.fullName if s.assigned_user else "ATANMADI"
            status_text = s.unassigned_reason if s.is_unassigned else (s.violation_warning or s.note or "-")

            table_data.append([
                Paragraph(s.date.strftime("%d.%m.%Y"), cell_style),
                Paragraph(day_name, cell_style),
                Paragraph(s.shift_type.name if s.shift_type else "Vardiya", cell_style),
                Paragraph(time_str, cell_style),
                Paragraph(f"{duration:.0f}s", cell_style),
                Paragraph(f"<b>{staff_name}</b>", cell_style),
                Paragraph(s.location or "-", cell_style),
                Paragraph(status_text, cell_style),
            ])

        # Widths for landscape A4 (Total width ~ 790pt)
        col_widths = [65, 65, 90, 80, 40, 160, 110, 170]
        schedule_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        schedule_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3157D5")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#3157D5")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        elements.append(schedule_table)

        # Approval Block
        elements.append(Spacer(1, 24))
        approval_data = [
            [
                Paragraph("<b>DÜZENLEYEN</b><br/><br/>Birim Sorumlusu / Yönetici<br/>İmza / Mühür", cell_style),
                Paragraph("<b>KONTROL EDEN</b><br/><br/>İdari İşler Sorumlusu<br/>İmza", cell_style),
                Paragraph("<b>ONAYLAYAN</b><br/><br/>Kurum Amiri / Başhekimlik<br/>Tarih / İmza / Onay", cell_style),
            ]
        ]
        approval_table = Table(approval_data, colWidths=[260, 260, 260])
        approval_table.setStyle(
            TableStyle([
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBEFORE", (1, 0), (1, -1), 0.5, colors.HexColor("#D1D5DB")),
                ("LINEBEFORE", (2, 0), (2, -1), 0.5, colors.HexColor("#D1D5DB")),
            ])
        )
        elements.append(KeepTogether([approval_table]))

        doc.build(elements)
        buffer.seek(0)
        return buffer
